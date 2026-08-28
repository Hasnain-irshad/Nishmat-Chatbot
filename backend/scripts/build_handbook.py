#!/usr/bin/env python3
"""
Build the Nishmat AI handbook as a Word document.

    python backend/scripts/build_handbook.py

Produces `docs/Nishmat AI Handbook.docx` — a downloadable, printable,
editable file to send to the client.

Word rather than PDF deliberately: the client works in Word (every one of her
131 lessons is a .docx), so she can open this, correct anything I have worded
badly, and send it on without needing another tool.

The content mirrors the web version at the artifact link. If you change one,
change the other.
"""

from __future__ import annotations

import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "Nishmat AI Handbook.docx"
LOGO = ROOT / "frontend" / "public" / "logo.jpg"

INK = RGBColor(0x1B, 0x17, 0x33)
MUTED = RGBColor(0x5A, 0x54, 0x7C)
LEARNER = RGBColor(0x4F, 0x3F, 0xBE)
ADMIN = RGBColor(0x8A, 0x64, 0x14)
SAFE = RGBColor(0x2F, 0x6D, 0x4F)


# --------------------------------------------------------------- content --

LEARNER_FEATURES = [
    ("Create an account and sign in",
     ["Anyone can sign up with an email address and password, or with their "
      "Google account. New accounts are always ordinary learner accounts — "
      "nobody can make themselves the teacher."],
     "Front page → Enter"),
    ("Browse all the lessons",
     ["Every published lesson is listed in order, with its Hebrew phrase and a "
      "short summary. There is a search box that looks through titles, Hebrew "
      "and lesson numbers, so “Moshia” or “#17” both find "
      "what you want."],
     "Lessons"),
    ("Read a lesson",
     ["The lesson appears the way it was written — the Hebrew phrase, how to say "
      "it, what it means, and then the lesson itself with its natural pauses and "
      "spacing kept exactly as intended.",
      "Hebrew reads right to left and keeps all its vowel marks."],
     "Lessons → any lesson"),
    ("Change how it reads",
     ["Two reading modes — a dark one for evening, and a light “parchment” "
      "one that looks like paper. Text size can be made bigger or smaller. Useful "
      "for reading on a phone."],
     "Top of any lesson"),
    ("Move through the series",
     ["Because each lesson follows on from the one before, every lesson has "
      "Previous and Next links at the bottom. You can read straight through the "
      "series in order."],
     "Bottom of any lesson"),
    ("Ask a question",
     ["Type any question and get an answer drawn from the lessons themselves — "
      "never from anywhere else. Each answer shows which lessons it came from, "
      "and those are clickable.",
      "If the lessons do not cover something, it says so plainly instead of "
      "guessing."],
     "Ask a question"),
    ("Come back to old conversations",
     ["Every conversation is saved and listed by when it happened — Today, "
      "Yesterday, This week. Open an old one and carry on where you stopped, or "
      "delete it.",
      "Conversations are completely private. Nobody else can see them, not even "
      "the teacher."],
     "Ask a question → side panel"),
]

ADMIN_FEATURES = [
    ("See everything at a glance",
     ["The dashboard shows how many lessons are published, how many are still "
      "drafts, how many need a second look, and how much has been spent on AI so "
      "far. Recently edited lessons are listed underneath."],
     "Dashboard"),
    ("Start a lesson by attaching a recording",
     ["Works like a messaging app. Press the + button and choose an audio "
      "recording, a document, or a photo — or simply type.",
      "No need to name it or number it. Both are worked out automatically and can "
      "be changed later."],
     "Create lesson"),
    ("Files are read automatically",
     ["A voice recording is turned into text. A Word file or PDF is read. A photo "
      "of a page is read too. Progress is shown while it happens.",
      "Hebrew and Yiddish words are handled specially, because ordinary "
      "transcription usually gets them wrong."],
     "Create lesson"),
    ("Correct the text before anything is written",
     ["A recording is never transcribed perfectly. The text can be corrected "
      "first, so a mistaken word never finds its way into the finished lesson."],
     "Lesson → source"),
    ("Find any lesson",
     ["All lessons in one list, with tabs for Published, Drafts, Awaiting publish "
      "and Needs review. Search works across titles, Hebrew and numbers.",
      "Labels show at a glance which lessons have unpublished edits waiting."],
     "All lessons"),
    ("Let the AI write the first draft",
     ["Press Generate draft and the lesson is written from the recording, in the "
      "teacher's own style and structure. It takes about a minute and shows what "
      "it is doing as it goes.",
      "The draft is never published on its own — it waits to be read."],
     "Lesson → Generate draft"),
    ("Edit it by hand",
     ["The lesson is shown in labelled parts — greeting, story, reflection, "
      "blessing — and each can be edited directly. Hebrew parts switch to "
      "right-to-left automatically.",
      "Word and line counts are shown for each part, because the short lines are "
      "part of the style."],
     "Lesson editor"),
    ("Ask the AI to change something",
     ["Hover over any part and press Ask AI. Say what you want in ordinary words "
      "— “make this warmer”, “this is too long”. Highlight "
      "just a paragraph first to change only that.",
      "The change is shown side by side with the original first. Nothing is kept "
      "until you say so."],
     "Lesson editor → Ask AI"),
    ("See it as a learner will",
     ["The Preview button shows the lesson exactly as a learner sees it, without "
      "leaving the editor. There is also “View as learner” at the top of "
      "every admin page."],
     "Lesson editor → Preview"),
    ("Go back to any earlier version",
     ["Every save is kept. The side panel lists them all, showing which was "
      "written by AI, which by hand, and which is currently live.",
      "Restoring an old version brings it back as a new one — nothing is ever "
      "lost or overwritten."],
     "Lesson editor → Version history"),
    ("Publish, or take back",
     ["A lesson only becomes visible to learners when Publish is pressed. It can "
      "be withdrawn again at any time, and it disappears from the "
      "question-answering too.",
      "Edits can be made to a published lesson without changing what learners "
      "see, until you publish the update."],
     "Lesson editor → Publish"),
    ("Change the lesson format and the writing style",
     ["The shape of a lesson — which parts it has, what order they come in — can "
      "be changed here, along with the written instructions that tell the AI how "
      "the teacher writes.",
      "Changes take effect on the next lesson straight away. No developer needed."],
     "Templates"),
]

SAFETY = [
    ("It never publishes by itself",
     ["No matter how good a draft is, it waits. Publishing is always a person "
      "pressing a button."]),
    ("It will not invent a source",
     ["The AI may only quote a pasuk, a Gemara, or a teaching that is actually in "
      "the recording. If it is not there, it is left out — even if the AI "
      "recognises it.",
      "Asked directly to “add a quote from the Gemara”, it refuses."]),
    ("It checks its own work",
     ["Every draft is read back by a second check that looks for invented "
      "material, altered Hebrew, missing parts, and writing that has drifted from "
      "the teacher's voice. Anything it finds is listed for you."]),
    ("Learners only ever see published work",
     ["Drafts and half-finished lessons are invisible to learners, and the "
      "question-answering will not use them either. The two always match."]),
    ("It cannot overspend",
     ["Every AI use is recorded and shown on the dashboard. There is a spending "
      "limit built in — once it is reached, the system stops rather than "
      "continuing to spend."]),
]

GLOSSARY = [
    ("Draft", "A lesson that has been written but not yet made visible to learners."),
    ("Published", "Live. Learners can read it, and it is used when answering questions."),
    ("Version", "A saved copy of the lesson at one moment. Every save makes a new "
                "one, and old ones are kept."),
    ("Source material", "The recording, document or photo the lesson was written from."),
    ("Template", "The shape of a lesson — which parts it has and how the writing "
                 "should sound."),
    ("Needs review", "A label meaning something looked unusual and a person should "
                     "check it."),
]


# ----------------------------------------------------------------- build --


def build() -> Path:
    doc = Document()

    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(1.1)
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = MUTED
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    # ---- cover -------------------------------------------------------
    if LOGO.exists():
        logo_para = doc.add_paragraph()
        logo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_para.add_run().add_picture(str(LOGO), width=Inches(1.15))

    _centred(doc, "Nishmat AI", size=30, bold=False, colour=INK, space_after=2)
    _centred(doc, "A  J O U R N E Y   O F   P R A I S E", size=8.5, colour=MUTED,
             space_after=18)
    _centred(doc, "Handbook", size=15, colour=ADMIN, space_after=14)
    _centred(
        doc,
        "What the website does, page by page, in plain language. There are two "
        "sides to it: the part everyone sees, and the part only the teacher sees.",
        size=10.5, colour=MUTED, space_after=20,
    )

    _who(doc, "Learners", LEARNER,
         "Anyone who signs up. They can read every published lesson and ask "
         "questions about them. They cannot change anything.")
    _who(doc, "The teacher (admin)", ADMIN,
         "Rivkah's account. She creates lessons, edits them, and decides when "
         "they go live. Only an admin account can reach these pages.")

    doc.add_page_break()

    # ---- the three parts ---------------------------------------------
    _heading(doc, "What learners can do", LEARNER, f"{len(LEARNER_FEATURES)} things")
    for title, paragraphs, where in LEARNER_FEATURES:
        _feature(doc, title, paragraphs, where, LEARNER)

    doc.add_page_break()

    _heading(doc, "What the teacher can do", ADMIN, f"{len(ADMIN_FEATURES)} things")
    for title, paragraphs, where in ADMIN_FEATURES:
        _feature(doc, title, paragraphs, where, ADMIN)

    doc.add_page_break()

    _heading(doc, "What the system will not do", SAFE, f"{len(SAFETY)} rules")
    for title, paragraphs in SAFETY:
        _feature(doc, title, paragraphs, None, SAFE)

    # ---- glossary -----------------------------------------------------
    _heading(doc, "A few words explained", MUTED, None)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for term, meaning in GLOSSARY:
        row = table.add_row().cells
        row[0].width = Inches(1.5)
        term_run = row[0].paragraphs[0].add_run(term)
        term_run.bold = True
        term_run.font.size = Pt(10)
        term_run.font.color.rgb = INK
        meaning_run = row[1].paragraphs[0].add_run(meaning)
        meaning_run.font.size = Pt(10)
        meaning_run.font.color.rgb = MUTED

    # ---- close --------------------------------------------------------
    doc.add_paragraph()
    closing = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hebrew = closing.add_run("נִשְׁמַת כָּל חַי תְּבָרֵךְ אֶת שִׁמְךָ")
    hebrew.font.size = Pt(14)
    hebrew.font.color.rgb = ADMIN
    hebrew.font.name = "David"

    _centred(doc, "“The soul of every living being shall bless Your Name.”",
             size=9, colour=MUTED, space_after=0)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


# ---------------------------------------------------------------- helpers --


def _centred(doc, text, *, size, colour, bold=False, space_after=8):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(space_after)
    run = para.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = colour
    run.bold = bold
    return para


def _who(doc, label, colour, description):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.left_indent = Inches(0.15)
    marker = para.add_run(f"{label}   ")
    marker.bold = True
    marker.font.size = Pt(10.5)
    marker.font.color.rgb = colour
    body = para.add_run(description)
    body.font.size = Pt(10)
    body.font.color.rgb = MUTED


def _heading(doc, text, colour, count):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(14)
    para.paragraph_format.space_after = Pt(10)
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = colour
    if count:
        tail = para.add_run(f"    {count}")
        tail.font.size = Pt(9)
        tail.font.color.rgb = MUTED


def _feature(doc, title, paragraphs, where, colour):
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(11)
    heading.paragraph_format.space_after = Pt(3)
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(11.5)
    run.font.color.rgb = INK

    for text in paragraphs:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.18)
        para.paragraph_format.space_after = Pt(4)
        body = para.add_run(text)
        body.font.size = Pt(10.5)
        body.font.color.rgb = MUTED

    if where:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.18)
        para.paragraph_format.space_after = Pt(2)
        chip = para.add_run(f"Where: {where}")
        chip.font.size = Pt(8.5)
        chip.bold = True
        chip.font.color.rgb = colour


if __name__ == "__main__":
    path = build()
    print(f"  Written: {path}")
    print(f"  {path.stat().st_size / 1024:.0f} KB")
